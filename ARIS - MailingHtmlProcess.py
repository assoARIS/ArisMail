# -*- coding: utf-8 -*-
"""
Created on Mon Feb 21 21:12:17 2022

@author: JMichel
"""

''' https://www.google.com/settings/security/lesssecureapps '''

from openpyxl import load_workbook
import smtplib, ssl, email, time, os, codecs

from email import encoders
from email.mime.base import MIMEBase
from email.mime.multipart import MIMEMultipart
from email.mime.text import MIMEText

''' variables '''
#
''' adresse mail & mot de passe expediteur '''
senderMail = 'anciensindosuez.aris@gmail.com'
senderPwd = 'Aris-1982*'
#
''' serveur mail & port d'envoi '''
smtpserver = "smtp.gmail.com"
port = 587
#
mailSubject = "ARIS - Séjour au Mont Dore du 03 au 10/09/2022"
mailMessage = ""
textMessage = ""
htmlMessage = ""
#
''' fichiers '''
fchXLS = "ARIS - MailingList - 20220311.xlsx"
feuilXLS = "TestPerso"
feuilLOG = "log"
feuilERR = "err"
fchTXT = ""
fchHTML = "ARIS - MailingWeb - 20220311.html"
fchPDF = "2022 LE MONT DORE présentation - 20220311.pdf"
#

''' fonctions '''
#
def clear():
    print("\033[H\033[J")
#
def setGlobal():
    global p,q,chk
    p , q = 0 , 0
    ''' variable de controle d'execution '''
    global chk
    chk = True
#
def showProcess(stp):
    global p
    p = p + 1
    print("-step:" +str(p) +" - " +stp +" -ok-")
#
def showDispatch(adr):
    global q
    q = q + 1
    print("-mail:" +str(q) +" - " +adr +" -ok-")
#
def readTextFile():
    from os.path import isfile
    try:
        if(isfile(fchTXT)):
            fchR = codecs.open(fchTXT,"r",encoding="utf-8")
            txt = fchR.read()
            fchR.close()
            return(txt)
        else:
            return("??")
    except:
        print("! erreur lecture fichier " +fchTXT +" !")
        global chk
        chk = False
#
def readHtmlFile():
    from os.path import isfile
    try:
        if(isfile(fchHTML)):
            fchR = codecs.open(fchHTML,"r",encoding="utf-8")
            html = fchR.read()
            fchR.close()
            return(html)
        else:
            return("??")
    except:
        print("! erreur lecture fichier " +fchHTML +" !")
        global chk
        chk = False
#
def getHtmlContent(prenom):
    html = "<!DOCTYPE html>"
    html += "<html lang='en'>"
    html += "<head>"
    html += "<meta charset='UTF-8'>"
    html += "<meta http-equiv='X-UA-Compatible' content='IE=edge'>"
    html += "<meta name='viewport' content='width=device-width, initial-scale=1.0'>"
    html += "<title>ARIS Info</title>"
    html += "</head"
    html += "<body>"
    html += "<div>cher(e) &nbsp;" + prenom + "</div>"
    html += "<br/>"
    html += readHtmlFile()
    html += "<br/><br/>"
    html += "<div style='font-size:x-small;text-align:center'>"
    html += "- mail non-commercial émanant de l'association <a href='http://www.anciensindosuez.org'>ARIS</a> et destiné à ses adhérents -"
    html += "<div/>"
    html += "</body>"
    html += "</html>"
    #
    return(html) 
#
def joinDocument():
    try:
        showProcess(fchPDF)
        ''' Open PDF file in binary mode '''
        with open(fchPDF, "rb") as attachment:
            ''' Add file as application/octet-stream
            client can usually download this automatically as attachment '''
            global pdfDoc
            pdfDoc = MIMEBase("application", "octet-stream")
            pdfDoc.set_payload(attachment.read())
            showProcess("pdfDoc")
            ''' Encode file in ASCII characters to send by email '''  
            encoders.encode_base64(pdfDoc)
            showProcess("pdfDoc encoding")
            ''' Add header as key/value pair to attached pdfDoc '''
            pdfDoc.add_header( "Content-Disposition"
                             , f"attachment; filename = {fchPDF}"
                             )
            showProcess("pdfDoc addHeader")
    except:
          print("! erreur attachement fichier pdf !")
          chk = False
          print("chk:",chk)
#
def startMailServer():
    ''' declaration du contexte de securite ssl '''
    sslContext = ssl.create_default_context()
    showProcess("sslContext")
    try:
        ''' lancement du serveur mail '''
        global mailServer
        mailServer = smtplib.SMTP(smtpserver,port)
        mailServer.ehlo()
        ''' cryptographie des envois '''
        mailServer.starttls(context = sslContext)
        mailServer.ehlo()
        ''' connexion au serveur mail '''
        mailServer.login(senderMail,senderPwd)
        showProcess("mailServer login")
    except:
        print("! erreur de lancement / connexion au serveur de mail !")
        chk = False
        print("-chk:",chk)
#
def openExcelFile(): 
    ''' ouverture du fichier Excel et selection du feuillet de données '''
    try:
        global workBook
        workBook = load_workbook(fchXLS)
        global workSheet
        workSheet = workBook[feuilXLS]
        showProcess("workBook & workSheet")
    except:
        print("! erreur de connexion au fichier Excel !")
        chk = False
        print("-chk:",chk)
#
def sendMailing():
   ''' lancement du processus d'execution d'envoi des mails à partir du fichier Excel ''' 
   print("\n")
   try:
       i = 2
       global workSheet
       while(workSheet["C"+str(i)].value != None):
           print("i:",i)
           if(workSheet["C"+str(i)].value != "??"
           and (workSheet["D"+str(i)].value == "R"
                or  workSheet["D"+str(i)].value == "CR")):
               ''' recuperation de l' adresse mail du destinataire '''
               receiverMail = workSheet["C"+str(i)].value
               ''' recuperation du et du prenom du destinataire '''
               nom = workSheet["A"+str(i)].value
               if(workSheet["B"+str(i)].value != None):
                   prenom = workSheet["B"+str(i)].value
               else:
                   prenom = ""
               #
               ''' redaction du contenu du mail '''
               ''' create a multipart message & set headers '''
               try:
                   mailHeader = MIMEMultipart()
                   mailHeader["From"] = senderMail
                   mailHeader["To"] = receiverMail
                   mailHeader["Subject"] = mailSubject
                   #
                   ''' add mailContent to email '''
                   mailHtml = getHtmlContent(prenom)
                   mailHeader.attach(MIMEText(mailHtml,"html"))
                   #
                   global mailContent
                   mailContent = mailHeader.as_string()
               except:
                   print("! erreur mailHeader : " + receiverMail + " !")
               #
               ''' envoi du mail '''
               try:
                   mailServer.sendmail(senderMail, receiverMail, mailContent)
                   showDispatch(receiverMail)
                   ''' temporisation entre envois '''
                   time.sleep(2)
               except:
                   print("! erreur envoi mail : " + receiverMail + " !")
               #
           i = i + 1
   except:
       err = str(i) + " - " + nom + " " + prenom + " : " + receiverMail
       print("! erreur processus envoi mailing : " + err + " !")
#
def closeExcelFile():
    ''' fermeture du fichier Excel '''
    print("\n")
    try:
        workBook.close()
        showProcess("workBook closing")
    except:
        print("! erreur de fermeture du fichier Excel !")
#
def stopMailServer():
    ''' arret du serveur de mail '''
    try:
        global mailServer
        mailServer.quit()
        showProcess("mailServer stop")
    except:
        print("! erreur de fermeture serveur mail !")
# 
    
    
''' execution '''
#
def main():
    clear()
    setGlobal()
    print("- start mailing -\n")
    if(chk == True):
        openExcelFile()
    if(chk == True):
        startMailServer()
    if(chk == True):
        sendMailing()
    closeExcelFile()
    stopMailServer()
    print("\n- end mailing -")
#   
if(__name__ == "__main__"):
  main()
#

